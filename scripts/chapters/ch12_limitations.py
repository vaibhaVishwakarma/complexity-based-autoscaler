#!/usr/bin/env python3
"""
word-doc/scripts/chapters/ch12_limitations.py: Chapter 12 Builder for BCE497J Project Report

Drafts Chapter 12: Limitations and Threats to Validity in a concise, rigorous format:
- Rulebook V2 typography (Times New Roman, 1.5 line spacing, 12pt body, justified).
- Academic Humanizer standards (neutral voice, varied sentence lengths, no em-dashes).
- Concise coverage of exchangeability bounds, systems constraints, and empirical validity threats.
- Zero inverted question marks, fully compliant OMML equations.
"""

from __future__ import annotations

import docx
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import nsdecls

from scripts.omml_builder import add_formatted_math_text, add_omml_equation


def build_chapter_12(doc: docx.Document):
    """Constructs a concise, rigorous academic appraisal for Chapter 12."""
    # Main Chapter Heading
    doc.add_heading("CHAPTER 12  LIMITATIONS AND THREATS TO VALIDITY", level=1)

    # -------------------------------------------------------------------------
    # 12.1 THEORETICAL ASSUMPTIONS AND EXCHANGEABILITY BOUNDS
    # -------------------------------------------------------------------------
    doc.add_heading("12.1 THEORETICAL ASSUMPTIONS AND EXCHANGEABILITY BOUNDS", level=2)

    p1 = doc.add_paragraph()
    p1.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    add_formatted_math_text(
        p1,
        "A rigorous scientific evaluation requires a candid appraisal of underlying theoretical assumptions. "
        "The edge triage mechanism in Conformal-Gated Dual-Horizon Autoscaling (CGDA) relies on Split Conformal Prediction, "
        "specifically Regularized Adaptive Prediction Sets (RAPS; Angelopoulos et al., 2021 [10]). The theoretical foundation "
        "of conformal prediction is exchangeability: calibration and test data are assumed to be identically distributed, or at least "
        "invariant under permutation. Under this exchangeability condition, conformal prediction provides a rigorous, "
        "distribution-free marginal coverage guarantee for any chosen miscoverage rate alpha in (0, 1):"
    )

    # OMML Equation 12.1: Marginal coverage guarantee
    ns_m = nsdecls("m")
    omml_eq12_1 = (
        f'<m:oMathPara {ns_m}>'
        '<m:oMath>'
        '<m:r><m:rPr><m:nor/></m:rPr><m:t>P(</m:t></m:r>'
        '<m:sSub>'
        '<m:e><m:r><m:rPr><m:sty m:val="i"/></m:rPr><m:t>Y</m:t></m:r></m:e>'
        '<m:sub><m:r><m:rPr><m:sty m:val="i"/></m:rPr><m:t>n</m:t></m:r><m:r><m:rPr><m:nor/></m:rPr><m:t>+1</m:t></m:r></m:sub>'
        '</m:sSub>'
        '<m:r><m:rPr><m:nor/></m:rPr><m:t> ∈ </m:t></m:r>'
        '<m:r><m:rPr><m:sty m:val="i"/></m:rPr><m:t>C</m:t></m:r>'
        '<m:r><m:t>(</m:t></m:r>'
        '<m:sSub>'
        '<m:e><m:r><m:rPr><m:sty m:val="i"/></m:rPr><m:t>X</m:t></m:r></m:e>'
        '<m:sub><m:r><m:rPr><m:sty m:val="i"/></m:rPr><m:t>n</m:t></m:r><m:r><m:rPr><m:nor/></m:rPr><m:t>+1</m:t></m:r></m:sub>'
        '</m:sSub>'
        '<m:r><m:rPr><m:nor/></m:rPr><m:t>)) ≥ 1 - </m:t></m:r>'
        '<m:r><m:rPr><m:sty m:val="i"/></m:rPr><m:t>α</m:t></m:r>'
        '</m:oMath>'
        '</m:oMathPara>'
    )
    add_omml_equation(doc, omml_eq12_1, eq_number="12.1")

    p2 = doc.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    add_formatted_math_text(
        p2,
        "In our system, calibrating RAPS on held-out Tiny ImageNet-C data with alpha = 0.10 yielded a non-conformity threshold "
        "q_hat = 0.79914, achieving 96.63% empirical validation coverage. In operational environments, however, exchangeability "
        "is subject to drift under sudden illumination changes, optical sensor corruptions, or unexpected domain shifts. When test inputs "
        "deviate sharply from the calibration distribution, marginal coverage guarantees can experience transient violations. "
        "Crucially, in the CGDA architecture, distribution drift acts favorably for systems stability: domain corruptions broaden prediction "
        "sets (|C(x)| >= 2), causing the edge gateway to shed ambiguous queries to the cloud and pre-warming GPU workers via the fast lookahead loop. "
        "Nonetheless, persistent non-stationarity necessitates periodic re-calibration or online adaptive conformal bounds (Gibbs & Candes, 2021 [13])."
    )

    # -------------------------------------------------------------------------
    # 12.2 SYSTEMS AND ACTUATION CONSTRAINTS
    # -------------------------------------------------------------------------
    doc.add_heading("12.2 SYSTEMS AND ACTUATION CONSTRAINTS", level=2)

    p3 = doc.add_paragraph()
    p3.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    add_formatted_math_text(
        p3,
        "The control performance of CGDA is fundamentally bounded by the physical dynamics of the cloud substrate. "
        "First, container startup latency establishes a hard lower bound on reactive response times. When the cold-start delay T_boot exceeds "
        "the SLA deadline D_SLA (15.0 s), no autoscaler operating without clairvoyant workload knowledge can prevent deadline violations. "
        "As established in Chapter 11, once T_boot > D_SLA, newly spawned containers cannot become operational before queuing delays miss the deadline. "
        "Second, the controller operates at a discrete actuation period of Delta t = 5.0 s. Traffic bursts with duration shorter than Delta t cannot "
        "be preempted during their initial arrival, relying entirely on existing standby capacity and edge buffering. "
        "Third, the asymmetric scale-down cooldown parameter tau_down = 15.0 s enforces stability at the expense of temporary over-provisioning: "
        "when demand suddenly subsides, excess workers are held active across three control epochs before being decommissioned."
    )

    # -------------------------------------------------------------------------
    # 12.3 THREATS TO EMPIRICAL VALIDITY
    # -------------------------------------------------------------------------
    doc.add_heading("12.3 THREATS TO EMPIRICAL VALIDITY", level=2)

    p4 = doc.add_paragraph()
    p4.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    add_formatted_math_text(
        p4,
        "To ensure objective scientific interpretation, we identify three potential threats to experimental validity. "
        "Internal validity is protected by executing all evaluations under fixed random seeds across matched synthetic and empirical Azure Functions "
        "workload traces, ensuring identical arrival patterns across compared autoscalers. A potential internal threat arises from network variance: "
        "wide-area latency was emulated using stochastic RTT distributions (15-35 ms) rather than live transcontinental links. "
        "Construct validity is maintained by evaluating standardized systems metrics: tail SLA miss percentage (p99 latency > 15.0 s), cumulative "
        "worker-seconds, actuator switching churn, and maximum queue backlog. "
        "External validity concerns the generalizability of findings beyond the specific experimental configuration. While empirical measurements "
        "were obtained on an AWS c7i serverless edge host (150 mCPU container quota) and NVIDIA Tesla T4 cloud GPUs running Triton Inference Server, "
        "the underlying two-loop control principles apply broadly to multi-tier cascading systems, including Vision-Language Models (VLMs) "
        "and cascading Large Language Models (LLMs) deployed across edge-cloud continuum topologies."
    )
